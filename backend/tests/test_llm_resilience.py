"""LLM integration: SSE streaming, session recovery, fallback, output parsing, documents."""
import io
import json

import docx
import pytest
import requests

import services.project_state_manager as state_manager
from models import AuditLog, Message, Project
from services.fdr_summary import FDR_SCHEMA, generate_fdr_json
from services.project_state_manager import clean_json_text, extract_delta_updates, get_structured_state
from tests.conftest import FakeResponse, sse
from tests.helpers import pending_fix, seed_messages


def events(response):
    """Parse the SSE body into a list of event dicts (skipping unparseable lines)."""
    parsed = []
    for line in response.text.splitlines():
        if line.startswith("data:"):
            try:
                parsed.append(json.loads(line[5:]))
            except ValueError:
                pass
    return parsed


def tokens(response):
    return "".join(e["data"] for e in events(response) if e.get("event") == "token")


def _fresh(db, project):
    db.expire_all()
    return db.get(Project, project.id)


@pytest.fixture
def chat(client, auth, make_user, make_project, llm):
    """A project owned by a user, plus a helper that sends a chat message to it."""
    owner = make_user()
    project = make_project(owner)

    class Chat:
        pass

    c = Chat()
    c.owner, c.project, c.llm = owner, project, llm
    c.send = lambda text="Tell me more": client.post(
        "/api/predict", headers=auth(owner), json={"question": text, "projectId": project.id}
    )
    return c


# ------------------------------------------------------------------ streaming happy path
class TestStreaming:
    def test_streams_tokens_as_server_sent_events(self, chat):
        r = chat.send()
        assert r.status_code == 200 and r.headers["content-type"].startswith("text/event-stream")
        assert tokens(r) == "".join(chat.llm.stream_tokens)
        assert any(e["event"] == "metadata" for e in events(r))

    def test_persists_both_sides_of_the_conversation(self, chat, db):
        chat.send("What do you need?")
        rows = db.query(Message).filter_by(project_id=chat.project.id).order_by(Message.id).all()
        assert [(m.role, m.text) for m in rows] == [("user", "What do you need?"), ("ai", "".join(chat.llm.stream_tokens).strip())]

    def test_sends_a_well_formed_request_to_the_provider(self, chat, db):
        chat.send("Budget is 10k")
        call = chat.llm.streaming_calls[0]
        session = _fresh(db, chat.project).forjinn_session_id
        assert call["payload"]["streaming"] is True
        assert call["payload"]["chatId"] and call["payload"]["overrideConfig"] == {"sessionId": call["payload"]["chatId"]}
        assert "Budget is 10k" in call["prompt"] and call["prompt"].rstrip().endswith("AI:")
        assert session  # the provider's session id was recorded for continuity

    def test_records_the_session_id_the_provider_returns(self, chat, db):
        chat.llm.stream_script = [chat.llm.stream_response(chat_id="session-from-provider")]
        chat.send()
        assert _fresh(db, chat.project).forjinn_session_id == "session-from-provider"

    def test_extracted_facts_land_in_project_state_after_the_stream(self, chat, db):
        chat.llm.delta_text = json.dumps({"department": "Logistics"})
        chat.send()
        assert get_structured_state(_fresh(db, chat.project))["department"] == "Logistics"

    def test_questions_the_ai_asked_are_remembered(self, chat, db):
        chat.send()
        assert "What is your budget?" in get_structured_state(_fresh(db, chat.project))["asked_questions"]

    def test_conversation_start_is_audited_once(self, chat, db):
        chat.send()
        chat.send()
        assert db.query(AuditLog).filter_by(action="conversation started").count() == 1

    def test_summarisation_kicks_in_at_ten_active_messages(self, chat, db):
        seed_messages(db, chat.project, 10)
        chat.llm.summary_text = "compacted"
        chat.send()
        assert _fresh(db, chat.project).summary == "compacted"
        assert db.query(Message).filter_by(project_id=chat.project.id, is_archived=True).count() > 0

    def test_state_extraction_failure_does_not_fail_the_chat(self, chat, db):
        chat.llm.nonstream_error = requests.ConnectionError("extractor down")
        r = chat.send()
        assert r.status_code == 200 and tokens(r)
        assert db.query(Message).filter_by(project_id=chat.project.id, role="ai").count() == 1


# ------------------------------------------------------------------ stream content handling
class TestStreamContent:
    def test_unrelated_events_are_forwarded_and_noise_is_dropped(self, chat):
        chat.llm.stream_script = [FakeResponse(200, lines=[
            b": keep-alive comment", b"", sse("start"), sse("token", "Hi."), b"data: {not json", sse("usage", {"n": 3}),
        ])]
        r = chat.send()
        assert r.status_code == 200
        assert ": keep-alive" not in r.text
        assert [e["event"] for e in events(r)] == ["start", "token", "usage"]
        assert "data: {not json" in r.text  # forwarded untouched, but never breaks parsing

    def test_a_reply_with_no_tokens_is_not_saved(self, chat, db):
        chat.llm.stream_script = [FakeResponse(200, lines=[sse("metadata", {"chatId": "x"})])]
        assert chat.send().status_code == 200
        assert db.query(Message).filter_by(project_id=chat.project.id, role="ai").count() == 0

    def test_non_session_errors_are_shown_to_the_client_not_retried(self, chat):
        chat.llm.stream_script = [FakeResponse(200, lines=[sse("error", message="rate limit exceeded")])]
        r = chat.send()
        assert any(e["event"] == "error" for e in events(r))
        assert len(chat.llm.streaming_calls) == 1


# ------------------------------------------------------------------ expired / missing provider session
class TestSessionRecovery:
    @pytest.mark.parametrize("status", [400, 404])
    def test_expired_session_response_recreates_the_session_and_retries_once(self, chat, db, status):
        original = "session-original"
        chat.project.forjinn_session_id = original
        db.commit()
        chat.llm.stream_script = [FakeResponse(status, text="session not found")]
        r = chat.send()

        assert r.status_code == 200 and tokens(r) == "".join(chat.llm.stream_tokens)
        first, second = chat.llm.streaming_calls
        assert first["payload"]["chatId"] == original
        assert second["payload"]["chatId"] not in (original, None)
        assert second["payload"]["overrideConfig"]["sessionId"] == second["payload"]["chatId"]
        fresh = _fresh(db, chat.project)
        assert fresh.forjinn_session_id != original  # the new session was persisted

    def test_in_stream_session_error_retries_and_discards_the_partial_reply(self, chat, db):
        chat.llm.stream_script = [FakeResponse(200, lines=[sse("token", "partial "), sse("error", message="Session not found")])]
        r = chat.send()

        assert len(chat.llm.streaming_calls) == 2
        assert tokens(r).endswith("".join(chat.llm.stream_tokens))
        assert not any(e["event"] == "error" for e in events(r))  # the error itself is never shown
        saved = db.query(Message).filter_by(project_id=chat.project.id, role="ai").one().text
        assert "partial" not in saved and saved == "".join(chat.llm.stream_tokens).strip()

    @pytest.mark.parametrize("expired", ["Session expired", "session not found"])
    def test_expiry_wording_is_matched_case_insensitively(self, chat, expired):
        chat.llm.stream_script = [FakeResponse(200, lines=[sse("error", message=expired)])]
        chat.send()
        assert len(chat.llm.streaming_calls) == 2

    def test_recovery_is_attempted_only_once(self, chat, db):
        chat.llm.stream_script = [FakeResponse(404, text="gone"), FakeResponse(404, text="gone")]
        r = chat.send()
        assert r.status_code == 200
        first, second = chat.llm.streaming_calls  # exactly two: never loops
        # Only the session actually used for the retry is persisted; a second failure must not mint another.
        assert _fresh(db, chat.project).forjinn_session_id == second["payload"]["chatId"] != first["payload"]["chatId"]


# ------------------------------------------------------------------ provider outage -> local fallback
class TestProviderOutage:
    @pytest.mark.parametrize(
        "failure",
        [requests.ConnectionError("refused"), requests.Timeout("timed out"), requests.HTTPError("503 Service Unavailable")],
        ids=["connection-refused", "timeout", "http-5xx"],
    )
    @pending_fix("issue 4", "the outage fallback makes an HTTP call back to this same server, which cannot succeed under load")
    def test_falls_back_to_a_local_response_without_any_http_call(self, chat, db, failure):
        chat.llm.stream_script = [failure]
        r = chat.send()  # the no-real-network guard makes any attempted HTTP call fail this test
        assert r.status_code == 200
        assert "mock response" in tokens(r)
        assert db.query(Message).filter_by(project_id=chat.project.id, role="ai").count() == 1


# ------------------------------------------------------------------ structured output parsing
class TestCleanJsonText:
    @pytest.mark.parametrize(
        "raw",
        [
            '{"a": 1}',
            '  {"a": 1}  \n',
            '```json\n{"a": 1}\n```',
            '```\n{"a": 1}\n```',
            '\n```json\n{"a": 1}\n```\n',
        ],
        ids=["plain", "padded", "json-fence", "bare-fence", "fence-with-outer-whitespace"],
    )
    def test_strips_markdown_wrappers(self, raw):
        assert json.loads(clean_json_text(raw)) == {"a": 1}

    def test_pretty_printed_json_inside_a_fence_survives(self):
        assert json.loads(clean_json_text('```json\n{\n  "a": [1, 2],\n  "b": {"c": "d"}\n}\n```')) == {"a": [1, 2], "b": {"c": "d"}}


class TestExtractDeltaUpdates:
    def _extract(self, llm, text):
        llm.delta_text = text
        return extract_delta_updates("user said", "ai said", dict(state_manager.DEFAULT_STATE))

    def test_well_formed_delta(self, llm):
        assert self._extract(llm, '{"department": "IT", "budget": "5k"}') == {"department": "IT", "budget": "5k"}

    def test_markdown_wrapped_delta(self, llm):
        assert self._extract(llm, '```json\n{"department": "IT"}\n```') == {"department": "IT"}

    def test_empty_delta(self, llm):
        assert self._extract(llm, "{}") == {}

    @pytest.mark.parametrize(
        "garbage",
        [
            "",
            "   ",
            "not json at all",
            'Sure! Here is the JSON: {"department": "IT"}',  # prose around it
            '{"department": "IT"',  # truncated
            "{'department': 'IT'}",  # single quotes
            '{"department": "IT",}',  # trailing comma
            '```json\n{"department": "IT"}```',  # closing fence not on its own line
            "null",
            "[]",
            '["department"]',
            "42",
            '"just a string"',
        ],
    )
    def test_unusable_output_yields_an_empty_delta_and_never_raises(self, llm, garbage):
        assert self._extract(llm, garbage) == {}

    @pytest.mark.parametrize(
        "body",
        [{"text": "{}"}, {"output": {"content": '{"department": "IT"}'}}, {"output": '{"department": "IT"}'}, {}, {"output": 7}],
        ids=["text", "output-dict", "output-str", "empty", "output-wrong-type"],
    )
    def test_provider_response_shapes(self, monkeypatch, body):
        monkeypatch.setattr(state_manager, "request_with_retry", lambda *a, **k: FakeResponse(200, json_body=body))
        result = extract_delta_updates("u", "a", dict(state_manager.DEFAULT_STATE))
        assert result == ({"department": "IT"} if "IT" in json.dumps(body) else {})

    @pytest.mark.parametrize("failure", [requests.ConnectionError("x"), requests.Timeout("x"), ValueError("bad body")])
    def test_provider_failures_yield_an_empty_delta(self, llm, failure):
        llm.nonstream_error = failure
        assert extract_delta_updates("u", "a", dict(state_manager.DEFAULT_STATE)) == {}

    def test_non_json_http_body_yields_an_empty_delta(self, monkeypatch):
        monkeypatch.setattr(state_manager, "request_with_retry", lambda *a, **k: FakeResponse(200, json_body=None, text="<html>"))
        assert extract_delta_updates("u", "a", dict(state_manager.DEFAULT_STATE)) == {}

    def test_the_current_state_and_exchange_are_given_to_the_model(self, llm):
        state = dict(state_manager.DEFAULT_STATE, project_name="Hub")
        extract_delta_updates("USER-TEXT", "AI-TEXT", state, active_section="Budget")
        prompt = llm.calls[-1]["prompt"]
        assert all(s in prompt for s in ("Hub", "USER-TEXT", "AI-TEXT", "Budget"))

    @pending_fix("issue 7", "the extracted delta is returned as-is, so a bare string for a list field reaches the state unvalidated")
    def test_a_bare_string_for_a_list_field_is_coerced(self, llm):
        result = self._extract(llm, '{"functional_requirements": "SSO required"}')
        assert result == {"functional_requirements": [{"title": "SSO required", "priority": "Medium", "confidence": 1.0}]}


# ------------------------------------------------------------------ document generation
def _docx_text(content):
    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


class TestFdrJson:
    @pytest.fixture
    def project(self, make_user, make_project):
        return make_project(make_user(), "Hub")

    def test_valid_response_is_merged_over_the_schema(self, llm, project):
        llm.fdr_text = json.dumps({"business_problem": "slow", "key_features": [{"feature": "Login", "description": "SSO"}]})
        result = generate_fdr_json(project)
        assert result["business_problem"] == "slow" and result["key_features"] == [{"feature": "Login", "description": "SSO"}]
        assert set(FDR_SCHEMA) <= set(result)

    def test_markdown_wrapped_response(self, llm, project):
        llm.fdr_text = '```json\n{"scope": "everything"}\n```'
        assert generate_fdr_json(project)["scope"] == "everything"

    def test_list_fields_with_the_wrong_shape_are_sanitised(self, llm, project):
        llm.fdr_text = json.dumps({"key_features": ["just a string", {"feature": "ok"}], "user_roles": "admin", "data_sources": [1, None]})
        result = generate_fdr_json(project)
        assert result["key_features"] == [{"feature": "ok", "description": ""}]
        assert result["user_roles"] == [] and result["data_sources"] == []

    @pytest.mark.parametrize("garbage", ["", "nonsense", "[1, 2]", "null", '{"scope": '])
    def test_unusable_response_yields_the_blank_schema(self, llm, project, garbage):
        llm.fdr_text = garbage
        assert generate_fdr_json(project) == FDR_SCHEMA

    def test_provider_failure_yields_the_blank_schema(self, llm, project):
        llm.nonstream_error = requests.ConnectionError("down")
        assert generate_fdr_json(project) == FDR_SCHEMA

    def test_the_whole_conversation_is_sent_in_order(self, db, llm, project):
        seed_messages(db, project, 3)
        generate_fdr_json(_fresh(db, project))
        prompt = llm.calls[-1]["prompt"]
        assert prompt.index("User: m0") < prompt.index("AI: m1") < prompt.index("User: m2")


class TestExport:
    @pytest.fixture
    def world(self, client, auth, make_user, make_project, llm):
        owner = make_user()
        project = make_project(owner, "Retail Hub")
        return owner, project, lambda fmt, who=None: client.get(f"/api/projects/{project.id}/export?format={fmt}", headers=auth(who or owner))

    def test_docx_is_a_real_document_containing_the_project(self, world, llm):
        llm.fdr_text = json.dumps({"business_problem": "Slow warehouse picking"})
        _, _, export = world
        r = export("docx")
        assert r.status_code == 200 and r.content[:2] == b"PK"
        assert "wordprocessingml" in r.headers["content-type"] and "Retail_Hub" in r.headers["content-disposition"]
        text = _docx_text(r.content)
        assert "Slow warehouse picking" in text and "Retail Hub" in text  # name falls back to the project's

    def test_undiscussed_fields_are_marked_missing_rather_than_left_blank(self, world, llm):
        llm.fdr_text = "{}"
        assert "[MISSING]" in _docx_text(world[2]("docx").content)

    def test_docx_is_still_produced_when_the_model_is_down(self, world, llm):
        llm.nonstream_error = requests.ConnectionError("down")
        r = world[2]("docx")
        assert r.status_code == 200 and "[MISSING]" in _docx_text(r.content)

    def test_pdf_is_a_real_document(self, world, llm):
        llm.markdown_text = "# Title\n\n## Section\n- point one\n1. numbered\nplain text with <tags> & symbols"
        r = world[2]("pdf")
        assert r.status_code == 200 and r.content.startswith(b"%PDF") and r.headers["content-type"] == "application/pdf"

    def test_pdf_is_still_produced_when_the_model_is_down(self, world, llm):
        llm.nonstream_error = requests.ConnectionError("down")
        r = world[2]("pdf")
        assert r.status_code == 200 and r.content.startswith(b"%PDF")

    def test_unsupported_format_is_rejected(self, world):
        assert world[2]("xlsx").status_code == 400

    def test_viewers_can_export_and_it_is_audited(self, client, db, world, make_user, add_member):
        _, project, export = world
        viewer = make_user()
        add_member(project, viewer)
        assert export("pdf", viewer).status_code == 200
        assert db.query(AuditLog).filter_by(action="document generation", project_id=project.id).count() == 1


# ------------------------------------------------------------------ the real retry wrapper
import types  # noqa: E402

import utils.prod_ready as prod_ready  # noqa: E402


class TestRequestWithRetry:
    """These use the real request_with_retry, with only the network and sleeping stubbed."""

    @pytest.fixture
    def net(self, monkeypatch):
        net = types.SimpleNamespace(calls=[], sleeps=[], script=[])

        def fake_request(method, url, **kwargs):
            net.calls.append({"method": method, "url": url, "kwargs": kwargs})
            outcome = net.script.pop(0) if net.script else FakeResponse(200, json_body={})
            if isinstance(outcome, Exception):
                raise outcome
            return outcome

        monkeypatch.setattr(prod_ready.requests, "request", fake_request)
        monkeypatch.setattr(prod_ready, "time", types.SimpleNamespace(sleep=net.sleeps.append))
        return net

    def test_success_needs_a_single_attempt(self, net):
        assert prod_ready.request_with_retry("POST", "http://x").status_code == 200
        assert len(net.calls) == 1 and net.sleeps == []

    def test_transient_failures_are_retried_with_exponential_backoff(self, net):
        net.script = [requests.ConnectionError("a"), requests.Timeout("b"), FakeResponse(200, json_body={"ok": 1})]
        assert prod_ready.request_with_retry("GET", "http://x").json() == {"ok": 1}
        assert len(net.calls) == 3 and net.sleeps == [1.0, 2.0]

    def test_gives_up_after_three_attempts_and_raises_the_last_error(self, net):
        net.script = [requests.ConnectionError("1"), requests.ConnectionError("2"), requests.ConnectionError("final")]
        with pytest.raises(requests.ConnectionError, match="final"):
            prod_ready.request_with_retry("GET", "http://x")
        assert len(net.calls) == 3 and net.sleeps == [1.0, 2.0]

    def test_timeouts_are_normalised_to_connect_and_read_pairs(self, net):
        prod_ready.request_with_retry("GET", "http://x")
        prod_ready.request_with_retry("GET", "http://x", timeout=30)
        prod_ready.request_with_retry("GET", "http://x", timeout=(1, 2))
        assert [c["kwargs"]["timeout"] for c in net.calls] == [(5, 90), (5, 30), (1, 2)]

    @pending_fix(
        "finding D",
        "4xx responses raise via raise_for_status and are then retried three times with backoff, "
        "although a client error can never succeed on retry",
    )
    def test_client_errors_are_not_retried(self, net):
        net.script = [FakeResponse(404), FakeResponse(404), FakeResponse(404)]
        with pytest.raises(requests.HTTPError):
            prod_ready.request_with_retry("POST", "http://x")
        assert len(net.calls) == 1


class TestExpiredSessionUnderRealConditions:
    @pending_fix(
        "finding D",
        "an expired provider session is reported as HTTP 404, but request_with_retry raises on 4xx, so the "
        "session-recovery branch (which expects a returned 404 response) is unreachable in production",
    )
    def test_http_404_for_an_expired_session_is_recovered(self, client, db, auth, make_user, make_project, monkeypatch):
        owner = make_user()
        project = make_project(owner)
        project.forjinn_session_id = "session-expired"
        db.commit()
        served = []

        def provider(method, url, **kwargs):
            body = kwargs.get("json") or {}
            served.append(body.get("chatId"))
            if not body.get("streaming"):
                return FakeResponse(200, json_body={"text": "{}"})
            if body.get("chatId") == "session-expired":  # a real provider rejects the dead session every time
                return FakeResponse(404, text="session not found")
            return FakeResponse(200, lines=[sse("token", "recovered")])

        monkeypatch.setattr(prod_ready.requests, "request", provider)
        monkeypatch.setattr(prod_ready, "time", types.SimpleNamespace(sleep=lambda s: None))
        r = client.post("/api/predict", headers=auth(owner), json={"question": "hi", "projectId": project.id})
        assert "recovered" in tokens(r)
        assert _fresh(db, project).forjinn_session_id != "session-expired"
