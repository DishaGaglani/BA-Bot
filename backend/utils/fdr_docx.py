import io
from docx import Document
from docx.shared import Pt

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)
MISSING = "[MISSING]"


def _run(paragraph, text, bold=False, size=FONT_SIZE):
    run = paragraph.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = size
    run.bold = bold
    return run


def _para(doc, text="", bold=False, size=FONT_SIZE):
    p = doc.add_paragraph()
    _run(p, text, bold=bold, size=size)
    return p


def _heading(doc, text):
    p = doc.add_paragraph()
    _run(p, text, bold=True, size=Pt(14))
    return p


def _subheading(doc, text):
    p = doc.add_paragraph()
    _run(p, text, bold=True, size=Pt(12))
    return p


def _set_cell(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    _run(p, text or MISSING, bold=bold)


def _val(data, key, default=MISSING):
    v = data.get(key)
    if v is None or v == "" or (isinstance(v, list) and not v):
        return default
    return v


def build_fdr_docx(data: dict) -> io.BytesIO:
    """Build the Requirement Discovery Form docx from a structured JSON dict,
    mirroring the layout of the reference 'Requirement Discover Format.docx' template."""
    doc = Document()

    # --- Project Information table ---
    info_table = doc.add_table(rows=2, cols=5)
    info_table.style = "Table Grid"
    headers = ["Project Name", "Department", "BU", "Project Sponsor", "Date"]
    keys = ["project_name", "department", "business_unit", "project_sponsor", "date"]
    for col, header in enumerate(headers):
        _set_cell(info_table.rows[0].cells[col], header, bold=True)
    for col, key in enumerate(keys):
        _set_cell(info_table.rows[1].cells[col], str(_val(data, key)))

    # --- Project Overview ---
    _heading(doc, "Project Overview")
    _subheading(doc, "Project Description")
    _para(doc, str(_val(data, "project_description")))
    _subheading(doc, "Stakeholders")
    stakeholders = data.get("stakeholders")
    if isinstance(stakeholders, list):
        stakeholders = ", ".join(str(s) for s in stakeholders) if stakeholders else None
    _para(doc, str(_val({"s": stakeholders}, "s")))

    # --- Discovery Phase ---
    _heading(doc, "Discovery Phase")
    _para(
        doc,
        "The Discovery phase is focused on understanding the broader business problem, "
        "user needs, and constraints, forming the foundation for the AI solution.",
    )
    discovery_table = doc.add_table(rows=4, cols=2)
    discovery_table.style = "Table Grid"
    discovery_rows = [
        ("Business Problem", "business_problem"),
        ("Business Goals", "business_goals"),
        ("Desired Outcomes", "desired_outcomes"),
        ("Other Insights", "other_insights"),
    ]
    for row, (label, key) in zip(discovery_table.rows, discovery_rows):
        _set_cell(row.cells[0], label, bold=True)
        _set_cell(row.cells[1], str(_val(data, key)))

    # --- High-Level Business Requirements ---
    _heading(doc, "High-Level Business Requirements")
    _subheading(doc, "3.1 Business Objectives")
    _para(doc, str(_val(data, "business_objectives")))
    _subheading(doc, "3.2 Existing Process and Pain Points")
    _para(doc, str(_val(data, "existing_process_pain_points")))
    _subheading(doc, "3.3 Success Criteria & Key Performance Indicators (KPIs)")
    _para(doc, str(_val(data, "success_criteria_kpis")))
    _subheading(doc, "3.4 Return on Investment Calculation")
    _para(doc, f"Operations impact (manhours saved): {_val(data, 'roi_operations_impact')}")
    _para(doc, f"Business impact (KPI): {_val(data, 'roi_business_impact')}")
    _subheading(doc, "3.5 Scope")
    _para(doc, str(_val(data, "scope")))

    # --- Functional Requirements ---
    _heading(doc, "Functional Requirements")
    _subheading(doc, "4.1 Key Features")
    features = data.get("key_features") or []
    feature_table = doc.add_table(rows=1 + max(len(features), 1), cols=2)
    feature_table.style = "Table Grid"
    _set_cell(feature_table.rows[0].cells[0], "Feature", bold=True)
    _set_cell(feature_table.rows[0].cells[1], "Description", bold=True)
    if features:
        for row, feat in zip(feature_table.rows[1:], features):
            _set_cell(row.cells[0], str(feat.get("feature", MISSING)))
            _set_cell(row.cells[1], str(feat.get("description", MISSING)))
    else:
        _set_cell(feature_table.rows[1].cells[0], MISSING)
        _set_cell(feature_table.rows[1].cells[1], MISSING)

    _subheading(doc, "4.2 User Roles & Permissions")
    roles = data.get("user_roles") or []
    role_table = doc.add_table(rows=1 + max(len(roles), 1), cols=2)
    role_table.style = "Table Grid"
    _set_cell(role_table.rows[0].cells[0], "Role", bold=True)
    _set_cell(role_table.rows[0].cells[1], "Permissions", bold=True)
    if roles:
        for row, role in zip(role_table.rows[1:], roles):
            _set_cell(row.cells[0], str(role.get("role", MISSING)))
            _set_cell(row.cells[1], str(role.get("permissions", MISSING)))
    else:
        _set_cell(role_table.rows[1].cells[0], MISSING)
        _set_cell(role_table.rows[1].cells[1], MISSING)

    # --- Non-Functional Requirements ---
    _heading(doc, "Non-Functional Requirements")
    _subheading(doc, "5.1 Performance")
    _para(doc, str(_val(data, "performance")))
    _subheading(doc, "5.2 Security")
    _para(doc, str(_val(data, "security")))
    _subheading(doc, "5.3 Scalability")
    _para(doc, str(_val(data, "scalability")))

    # --- System Requirements ---
    _heading(doc, "System Requirements")
    _subheading(doc, "6.1 Platforms/Devices")
    _para(doc, str(_val(data, "platforms_devices")))
    _subheading(doc, "6.2 Integration")
    _para(doc, str(_val(data, "integration")))
    _subheading(doc, "6.3 Data Requirements")
    _para(doc, str(_val(data, "data_requirements")))

    # --- Data Availability ---
    _heading(doc, "Data Availability")
    _para(doc, f"Nature of the data (Structured / Unstructured): {_val(data, 'nature_of_data')}")
    _subheading(doc, "7.1 Data Source & Location")
    sources = data.get("data_sources") or []
    source_table = doc.add_table(rows=1 + max(len(sources), 1), cols=3)
    source_table.style = "Table Grid"
    for col, header in enumerate(["Data", "Source", "Location"]):
        _set_cell(source_table.rows[0].cells[col], header, bold=True)
    if sources:
        for row, src in zip(source_table.rows[1:], sources):
            _set_cell(row.cells[0], str(src.get("data", MISSING)))
            _set_cell(row.cells[1], str(src.get("source", MISSING)))
            _set_cell(row.cells[2], str(src.get("location", MISSING)))
    else:
        for col in range(3):
            _set_cell(source_table.rows[1].cells[col], MISSING)

    # --- Approval & Sign-Off ---
    _heading(doc, "Approval & Sign-Off")
    _subheading(doc, "Project Manager")
    _para(doc, f"Name: {_val(data, 'project_manager_name')}")
    _para(doc, f"Date: {_val(data, 'project_manager_date')}")
    _subheading(doc, "Sponsors")
    _para(doc, f"Names and Signatures: {_val(data, 'sponsors_names')}")
    _para(doc, f"Date: {_val(data, 'sponsors_date')}")

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream
